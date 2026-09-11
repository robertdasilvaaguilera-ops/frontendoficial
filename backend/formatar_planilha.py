"""
Aplica formatacao profissional na planilha (usando openpyxl, seguindo
boas praticas: fonte legivel, cabecalho destacado com a cor da marca,
moeda de verdade na coluna de economia_estimada, colunas com largura
ajustada).

Roda uma vez, formata o arquivo que ja existe. Pode rodar de novo a
qualquer momento (idempotente - so reaplica o estilo).
"""
import os
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

PASTA_BASE = os.path.dirname(os.path.abspath(__file__))
CAMINHO_EXCEL = os.path.join(PASTA_BASE, "reports", "oportunidades.xlsx")

COR_NAVY = "1B2A4A"
COR_BRANCO = "FFFFFF"

# Colunas que devem ser formatadas como moeda (R$)
COLUNAS_MOEDA = {"economia_estimada"}


def main():
    if not os.path.exists(CAMINHO_EXCEL):
        print(f"[ERRO] Planilha não encontrada: {CAMINHO_EXCEL}")
        return

    wb = load_workbook(CAMINHO_EXCEL)
    ws = wb.active

    fonte_padrao = Font(name="Arial", size=10)
    fonte_cabecalho = Font(name="Arial", size=10, bold=True, color=COR_BRANCO)
    preenchimento_cabecalho = PatternFill(start_color=COR_NAVY, end_color=COR_NAVY, fill_type="solid")

    cabecalhos = {}
    for indice_coluna, celula in enumerate(ws[1], start=1):
        celula.font = fonte_cabecalho
        celula.fill = preenchimento_cabecalho
        celula.alignment = Alignment(horizontal="center", vertical="center")
        cabecalhos[celula.value] = indice_coluna

    for linha in ws.iter_rows(min_row=2):
        for celula in linha:
            celula.font = fonte_padrao

    for nome_coluna in COLUNAS_MOEDA:
        if nome_coluna in cabecalhos:
            indice = cabecalhos[nome_coluna]
            letra = get_column_letter(indice)
            for linha_num in range(2, ws.max_row + 1):
                celula = ws[f"{letra}{linha_num}"]
                celula.number_format = 'R$ #,##0.00;(R$ #,##0.00);"-"'

    for indice_coluna in range(1, ws.max_column + 1):
        letra = get_column_letter(indice_coluna)
        maior_largura = 0
        for celula in ws[letra]:
            valor = str(celula.value) if celula.value is not None else ""
            maior_largura = max(maior_largura, len(valor))
        ws.column_dimensions[letra].width = min(maior_largura + 2, 60)

    ws.freeze_panes = "A2"  # cabecalho fixo ao rolar

    wb.save(CAMINHO_EXCEL)
    print(f"Planilha formatada: {CAMINHO_EXCEL}")


if __name__ == "__main__":
    main()
